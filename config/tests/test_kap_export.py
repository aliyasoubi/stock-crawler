import base64
import json
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from stock_crawler.kap import SourceError, UnknownTicker
from stock_crawler.kap_export import (CompanyRegistry, ExportEntry, KapExportClient, archive_export, export_payload,
    load_manifest, parse_export_row, read_export, read_export_blob)
from stock_crawler.storage import RawStore, StateStore, StorageError
from stock_crawler.fetch import PacedClient
from stock_crawler.pipeline import Pipeline, SyncOptions
from stock_crawler.parser import PARSER_VERSION

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / 'tests/fixtures/kap/exports/two_companies_2023_2024.xlsx'
REGISTRY = ROOT / 'config/kap_companies.json'


def modified_book(change):
    import warnings
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        w = load_workbook(BytesIO(BOOK.read_bytes()))
    change(w.active)
    out = BytesIO()
    w.save(out)
    w.close()
    return out.getvalue()


def test_real_export_four_rows_and_source_units():
    rows = read_export(BOOK.read_bytes())
    registry = CompanyRegistry(REGISTRY)
    reports = {(registry.match(r['Company']).ticker, int(r['Year'])): parse_export_row(r, calendar_year_confirmed=True, parser_version=PARSER_VERSION) for r in rows}
    assert len(reports) == 4 and all(r.parse_status.value == 'valid' for r in reports.values())
    thyao = reports['THYAO', 2024].current_period()
    asels = reports['ASELS', 2024].current_period()
    assert thyao.total_assets == Decimal('1399606000000')
    assert thyao.revenue == Decimal('745430000000')
    assert thyao.profit_attributable_to_non_controlling_interests == Decimal('-21000000')
    assert asels.revenue == Decimal('120205594000')
    assert (thyao.currency_scale, asels.currency_scale) == (1000000, 1000)
    assert thyao.ebitda is None and thyao.finance_sector_revenue is None
    assert str(thyao.measuring_unit_date) == '2024-12-31'


def test_non_calendar_issuer_and_unreviewed_family_are_not_silently_published():
    row = read_export(BOOK.read_bytes())[0]
    assert parse_export_row(row, calendar_year_confirmed=False, parser_version=PARSER_VERSION).parse_status.value == 'unsupported'
    assert parse_export_row({**row, 'Sectoral Statement Type': 'bank'}, calendar_year_confirmed=True, parser_version=PARSER_VERSION).parse_status.value == 'unsupported'


@pytest.mark.parametrize('data', [b'<html>Access denied</html>' * 100, b'{"error":"denied"}', b'PK\x03\x04' + b'a' * 200])
def test_error_payload_never_becomes_success(data):
    with pytest.raises(SourceError):
        read_export(data)


def test_formula_and_unknown_column_rejected():
    with pytest.raises(SourceError, match='formula'):
        read_export(modified_book(lambda s: setattr(s['I8'], 'value', '=1+1')))
    with pytest.raises(SourceError, match='unknown export header'):
        read_export(modified_book(lambda s: setattr(s['I7'], 'value', 'Some other concept')))


def test_legacy_manifest_import_is_deduplicated_and_captures_unknown_date(tmp_path, clock):
    raw = RawStore(tmp_path)
    batch = {'batchIndex': 0, 'tickers': ['THYAO', 'ASELS'], 'base64': base64.b64encode(BOOK.read_bytes()).decode()}
    path = tmp_path / 'manifest.json'
    path.write_text(json.dumps([batch, batch]))
    entries, errors = load_manifest(path, CompanyRegistry(REGISTRY), raw, clock=clock)
    assert not errors and len(entries) == 4
    assert all(e.retrieved_at == clock() and not e.retrieved_at_known for e in entries)
    assert len(list((raw.root / '_exports').glob('*/source.xlsx'))) == 1
    assert read_export_blob(raw, entries[0].blob_hash) == BOOK.read_bytes()


def test_bad_batch_does_not_hide_good_batches(tmp_path):
    batch = {'tickers': ['THYAO', 'ASELS'], 'base64': base64.b64encode(BOOK.read_bytes()).decode()}
    path = tmp_path / 'manifest.json'
    path.write_text(json.dumps({'schemaVersion': 6, 'batches': [{**batch, 'sha256': 'bad'}, batch]}))
    entries, errors = load_manifest(path, CompanyRegistry(REGISTRY), RawStore(tmp_path))
    assert len(entries) == 4 and len(errors) == 1 and 'SHA-256' in errors[0]['error']


def test_company_identity_never_inferred_from_batch_order():
    r = CompanyRegistry(REGISTRY)
    with pytest.raises(UnknownTicker):
        r.match('UNVERIFIED HISTORICAL NAME', ['THYAO'])
    with pytest.raises(UnknownTicker):
        r.match(r.by_ticker['ASELS']['company_name'], ['THYAO'])
    assert len(r.entries) == 754


def _mock_export_client(settings, clock, data, requests):
    def respond(request):
        requests.append(request)
        return httpx.Response(200, content=data, headers={'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})
    settings.kap_company_registry = REGISTRY
    settings.kap_years = [2023, 2024]
    state = StateStore(settings.data_dir, clock=clock)
    fetcher = PacedClient(httpx.Client(transport=httpx.MockTransport(respond)), settings, state, allowed_hosts={'www.kap.org.tr'}, clock=clock)
    return KapExportClient(fetcher, settings, clock=clock), state, fetcher


def test_live_export_source_and_pipeline_use_same_real_workbook(settings, repo, clock):
    # Transport is mocked; response rows are unchanged from the real sample (THYAO 2023 + 2024 remain).
    data = modified_book(lambda s: s.delete_rows(8, 2))
    requests = []
    source, state, fetcher = _mock_export_client(settings, clock, data, requests)
    p = Pipeline(settings, repo, RawStore(settings.data_dir), state, source, clock=clock)
    summary = p.sync(SyncOptions(tickers=['THYAO']))
    rows = summary.data['companies']
    assert [(r['fiscal_year'], r['status']) for r in rows] == [(2023, 'published'), (2024, 'published')]
    assert len(requests) == 1 and requests[0].method == 'POST'
    request = json.loads(requests[0].content)
    assert request['yearList'] == ['2023', '2024'] and len(request['itemIdList']) == 10
    company = repo.get_company_by_ticker('kap_compare', 'THYAO')
    assert repo.current_view_row(company.company_id, 2024, 'consolidated')['revenue'] == Decimal('745430000000')
    assert repo.current_view_row(company.company_id, 2023, 'consolidated')['revenue'] == Decimal('504398000000')
    again = p.sync(SyncOptions(tickers=['THYAO'], refresh=True))
    assert all(r['status'] == 'already_parsed' for r in again.data['companies']) and len(repo.reports) == 2
    assert len(requests) == 2, 'refresh must request new bytes when the client is reused'
    # A new parser re-reads verified native bytes entirely offline.
    requests.clear()
    p.parser_version = '1.2.0'
    assert all(r['status'] == 'published' for r in p.reprocess(['THYAO']).data['companies'])
    assert not requests
    fetcher.close()


def test_sync_batches_companies_into_one_export_request(settings, repo, clock):
    requests = []
    source, state, fetcher = _mock_export_client(settings, clock, BOOK.read_bytes(), requests)
    p = Pipeline(settings, repo, RawStore(settings.data_dir), state, source, clock=clock)
    summary = p.sync(SyncOptions(tickers=['ASELS', 'THYAO']))
    assert len(requests) == 1, 'both companies share one POST'
    payload = json.loads(requests[0].content)
    assert len(payload['mkkMemberIdList']) == 2
    assert sorted((r['ticker'], r['fiscal_year']) for r in summary.data['companies']) == [('ASELS', 2023), ('ASELS', 2024), ('THYAO', 2023), ('THYAO', 2024)]
    assert len(repo.reports) == 4
    fetcher.close()


def test_batch_failure_is_reported_once_per_member_without_retrying(settings, repo, clock):
    calls = []
    def respond(request):
        calls.append(request)
        return httpx.Response(500, content=b'boom')
    settings.kap_company_registry = REGISTRY
    settings.kap_years = [2024]
    settings.max_retries = 0
    state = StateStore(settings.data_dir, clock=clock)
    fetcher = PacedClient(httpx.Client(transport=httpx.MockTransport(respond)), settings, state, allowed_hosts={'www.kap.org.tr'}, clock=clock)
    source = KapExportClient(fetcher, settings, clock=clock)
    p = Pipeline(settings, repo, RawStore(settings.data_dir), state, source, clock=clock)
    summary = p.sync(SyncOptions(tickers=['ASELS', 'THYAO']))
    assert len(calls) == 1
    assert [r['status'] for r in summary.data['companies']] == ['error', 'error']
    fetcher.close()


def test_non_calendar_issuer_is_skipped_before_any_request(settings, repo, clock):
    requests = []
    source, state, fetcher = _mock_export_client(settings, clock, BOOK.read_bytes(), requests)
    settings.kap_non_calendar_year_tickers = ['THYAO']
    p = Pipeline(settings, repo, RawStore(settings.data_dir), state, source, clock=clock)
    summary = p.sync(SyncOptions(tickers=['THYAO']))
    assert summary.data['companies'][0]['status'] == 'error' and 'KAP_NON_CALENDAR_YEAR_TICKERS' in summary.data['companies'][0]['error']
    assert not requests
    fetcher.close()


def test_historical_import_publishes_all_periods_and_reprocesses_all(settings, repo, clock):
    raw = RawStore(settings.data_dir)
    digest = archive_export(raw, BOOK.read_bytes())
    registry = CompanyRegistry(REGISTRY)
    entries = [ExportEntry(registry.match(r['Company']), r, digest, clock()) for r in read_export(BOOK.read_bytes())]
    from types import SimpleNamespace
    p = Pipeline(settings, repo, raw, StateStore(settings.data_dir), SimpleNamespace(market_source='kap_compare'), clock=clock)
    assert len(p.ingest_export_entries(entries).data['companies']) == 4
    assert len(repo.reports) == 4
    assert all(c['status'] == 'already_parsed' for c in p.ingest_export_entries(entries).data['companies'])
    repo.mark_withdrawn('kap_compare', '1396940')
    p.parser_version = '1.2.0'
    results = p.reprocess(['THYAO', 'ASELS']).data['companies']
    assert len(results) == 4 and sum(c['status'] == 'withdrawn' for c in results) == 1
    assert len(repo.reports) == 7
    # Raw-byte corruption prevents a further parser version from being published.
    (raw.root / '_exports' / digest / 'source.xlsx').write_bytes(b'corrupt')
    p.parser_version = '1.3.0'
    results = p.reprocess(['ASELS']).data['companies']
    assert all(c['status'] == 'cache_miss' for c in results)
    assert len(repo.reports) == 7


def test_identical_rows_in_repacked_workbook_have_same_snapshot(settings, repo, clock):
    from types import SimpleNamespace
    raw = RawStore(settings.data_dir)
    registry = CompanyRegistry(REGISTRY)
    p = Pipeline(settings, repo, raw, StateStore(settings.data_dir), SimpleNamespace(market_source='kap_compare'), clock=clock)
    for data in [BOOK.read_bytes(), modified_book(lambda s: None)]:
        digest = archive_export(raw, data)
        rows = [ExportEntry(registry.match(r['Company']), r, digest, clock()) for r in read_export(data)]
        p.ingest_export_entries(rows)
    assert len(repo.reports) == 4


def test_unmatched_row_is_rejected_without_discarding_the_rest_of_the_batch(settings, repo, clock):
    """Regression: the per-row handler once called an undefined `log`, so the resulting
    NameError reached the batch-level handler and failed every company in the POST."""
    data = modified_book(lambda s: s.cell(row=8, column=1, value='SOME RENAMED COMPANY A.Ş.'))
    requests = []
    source, state, fetcher = _mock_export_client(settings, clock, data, requests)
    p = Pipeline(settings, repo, RawStore(settings.data_dir), state, source, clock=clock)
    summary = p.sync(SyncOptions(tickers=['ASELS', 'THYAO']))
    published = {(r['ticker'], r['fiscal_year']) for r in summary.data['companies'] if r['status'] == 'published'}
    assert published, 'the surviving rows must still publish'
    assert summary.data['rejected_row_count'] == 1
    assert 'SOME RENAMED COMPANY' in summary.data['rejected_rows'][0]['company']
    fetcher.close()


def test_changing_kap_years_makes_a_recently_synced_company_due_again(settings, repo, clock):
    """Freshness must consider which years were collected, not only when the company was last
    touched; otherwise a historical backfill with a new KAP_YEARS collects nothing."""
    requests = []
    source, state, fetcher = _mock_export_client(settings, clock, BOOK.read_bytes(), requests)
    p = Pipeline(settings, repo, RawStore(settings.data_dir), state, source, clock=clock)
    assert any(r['status'] == 'published' for r in p.sync(SyncOptions(tickers=['THYAO'])).data['companies'])
    assert all(r['status'] == 'fresh' for r in p.sync(SyncOptions(tickers=['THYAO'])).data['companies'])
    settings.kap_years = [2019, 2020]
    assert not any(r['status'] == 'fresh' for r in p.sync(SyncOptions(tickers=['THYAO'])).data['companies'])
    fetcher.close()
